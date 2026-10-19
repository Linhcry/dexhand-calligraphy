"""Editable support documents; official signed forms must come from registration system."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables/报名准备_20260910'
CASE=ROOT/'outputs/blind_competition/20260910_130143_ac00cb'

def base(title):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(4)
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Title'].paragraph_format.space_before=Pt(0)
    doc.styles['Title'].paragraph_format.space_after=Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.add_paragraph(title,'Title')
    sec.footer.paragraphs[0].text='五指灵巧手比赛准备资料  |  2026年9月10日  |  参考仿真'
    sec.footer.paragraphs[0].runs[0].font.size=Pt(9)
    return doc

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for i,(h,width) in enumerate(zip(headers,widths)):t.columns[i].width=Inches(width);t.rows[0].cells[i].text=h
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EAF0F5');pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                for run in p.runs:run.font.size=Pt(10);run.bold=ri==0
    doc.add_paragraph()
    return t

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    r=json.loads((CASE/'run/report.json').read_text(encoding='utf-8'))
    doc=base('五指灵巧手技术资格仿真说明')
    doc.add_paragraph('本材料供技术组了解当前软件与参考仿真实验。已完成示例诗题的离线视觉识别、播报、持笔、蘸墨和全文书写记录；尚未接入本届官方设备，不能作为真机验收或官方评分证明。仿真材料是否适用于资格审核，待技术组确认。')
    doc.add_heading('队伍与项目',level=1)
    doc.add_paragraph('学校全称：________________    队伍名称：________________\n队长姓名：________________    指导老师：________________\n报名编号：________________    联系方式：________________')
    doc.add_heading('系统与自主流程',level=1)
    doc.add_paragraph('参考模型为DexHand021右手与JAKA Zu7，来自DexRobot官方开源仓库。使用本地MuJoCo、RapidOCR及笔顺素材。抽题程序独立生成屏幕图像，控制器在OCR后查询本地示例诗库，再播报并执行抓笔、连续持稳、蘸墨、落笔和抬笔。运行不依赖现场人工远程操作。')
    doc.add_paragraph('独立模块包括RGB-D几何感知、标定校验、纸面坐标排版、回合状态机与设备接口。新场景感知组件尚未全部接入物理执行闭环；当前物理执行仍使用仿真笔尖与接触力反馈。')
    doc.add_heading('实测过程指标',level=1)
    rows=[('视觉识别诗题',r['recognized_title']),('完成笔画路径',f"{r['completed_strokes']} / {r['planned_strokes']}"),
          ('完整仿真耗时',f"{r['elapsed_sim_s']:.2f}秒 / 限时300秒"),('连续稳定持笔',f"最长{r['longest_hold_s']:.2f}秒"),
          ('最大浸墨深度及停留',f"{r['dip_maximum_depth_mm']:.2f}毫米；{r['dip_longest_dwell_s']:.2f}秒"),
          ('移动机械臂与掌部最小离桌',f"{r['minimum_moving_arm_table_clearance_mm']:.2f}毫米"),
          ('运笔阶段有墨接触比例',f"{r['writing_ink_contact_fraction']*100:.2f}%"),
          ('含恢复等待的落笔接触比例',f"{r['all_pen_down_contact_fraction']*100:.2f}%"),
          ('轨迹XY误差中位数',f"{r['xy_error_median_mm']:.2f}毫米")]
    table(doc,['指标','结果'],rows,[3.1,3.7])
    doc.add_paragraph('复核记录：20260910_130143_ac00cb，9项过程检查通过，未计算官方分数。官方依据为第一轮通知第7–8页、五指规则第7–15及19–20页，原件附后。')
    doc.add_paragraph('补充复测：连续运笔方案于本日再次随机抽到《登鹳雀楼》，265.54秒完成131笔，9项检查通过。该结果仅支持同题重复性，不代表跨诗题稳健性。')
    doc.add_page_break();doc.add_heading('实际墨迹与验证范围',level=1)
    doc.add_picture(str(CASE/'run/result.png'),width=Inches(6.0))
    doc.paragraphs[-1].alignment=1
    doc.add_paragraph('图中右侧为由真实模拟接触记录生成的墨迹，不是字体贴图。笔画路径完成不代表书法质量达标，仍可见断墨与笔画粗细不均；没有官方字帖相似度或正式得分。固定底座网格最低离桌约0.85毫米，不能将移动臂指标误解为所有部件的最小距离。')
    doc.add_heading('视频说明',level=1)
    doc.add_paragraph('随附短视频选取本次完整物理记录的识别、抓取、蘸墨、运笔及完成片段，保持1倍仿真速度，明确标注片段剪辑和参考仿真。未通过加速视频将完整用时伪装为2分钟以内。视频时长、大小、源记录与剪辑时间段见随附验证JSON。完整状态记录和未剪辑回放可用于复核。')
    doc.add_heading('重复试验与局限',level=1)
    doc.add_paragraph('同题连续轨迹版本耗时265.18秒，匀速版本232.02秒；但XY误差中位数从0.79增至1.33毫米。匀速版本随后随机抽到《春晓》，在206.84秒因笔杆倾斜停止，仅完成112/149笔。增大夹持力的另一轮在52.93秒丢失握持。故当前展示是成功案例，不能代表所有诗题稳定成功；默认参数已恢复原基线，提速参数仅供实验。')
    doc.add_paragraph('机械臂型号、手的电机耦合、控制协议、真实触觉、工具变换、材料摩擦、笔尖柔顺和墨水模型尚未在官方设备验证。正式题库、字帖和龙图库尚未取得。26种RGB-D场景属于感知与规划验证，不是26次完整抓笔书写成功。')
    doc.add_paragraph('取得SDK后先验证只读反馈、单位和坐标，完成驱动看门狗及急停、碰撞模型和工具标定，再依次开展离桌路径、抓取、持稳、蘸墨、单字及整首的低速真机验证。当前真实设备接口未配置时拒绝启动。')
    doc.save(OUT/'技术资格仿真说明.docx')
    doc=base('五指赛项队伍信息采集表')
    doc.add_paragraph('用于整理填报信息，不替代官方报名系统生成的报名表或免责声明。请按实际情况填写；签字、学院盖章与扫描PDF应在正式表格上办理。')
    for label in ['学校全称','二级学院','队伍名称','报名系统队伍编号','队长姓名与电话','联系邮箱','校赛负责人及联系方式','指导老师姓名与联系方式']:
        doc.add_paragraph(label+'：________________________________________________')
    doc.add_heading('参赛队员',level=1)
    table(doc,['序号','姓名','学号','专业与年级'],[(i,'','','') for i in range(1,9)],[.5,1.4,2.,2.9])
    doc.add_paragraph('队员3至8人，指导老师1至3人，实际名单以赛项规则与报名系统核验为准。请勿将未确认成员填入正式提交材料。')
    doc.add_paragraph('办理状态：□ 系统填报  □ 负责人已提交晋级名单  □ 正式报名表已下载\n□ 全体成员及老师签字  □ 二级学院盖章  □ 扫描PDF上传  □ 技术资格确认')
    doc.save(OUT/'队伍信息采集表.docx')
    print(OUT)
if __name__=='__main__':main()
